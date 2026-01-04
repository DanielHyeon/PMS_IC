"""
개선된 RAG (Retrieval-Augmented Generation) 서비스
MinerU2.5 기반 Layout-Aware Document Parsing 적용
"""

import chromadb
from sentence_transformers import SentenceTransformer
from typing import List, Dict, Optional
import logging
import os
from pypdf import PdfReader
from docx import Document
import openpyxl

from document_parser import parse_and_chunk_document, MinerUDocumentParser, LayoutAwareChunker

logger = logging.getLogger(__name__)


class RAGServiceV2:
    """
    MinerU2.5 기반 개선된 RAG 서비스

    주요 개선사항:
    1. Layout-Aware Chunking: 문서 구조 기반 의미 단위 청킹
    2. 구조 정보 보존: 제목, 표, 리스트 등의 구조 정보 유지
    3. 컨텍스트 강화: 각 청크에 제목/섹션 정보 포함
    4. 메타데이터 활용: 구조 타입별 검색 가능
    """

    def __init__(self):
        # ChromaDB 클라이언트 초기화
        chroma_host = os.getenv("CHROMA_HOST", "localhost")
        chroma_port = int(os.getenv("CHROMA_PORT", "8000"))

        logger.info(f"Connecting to ChromaDB at {chroma_host}:{chroma_port}")
        self.chroma_client = chromadb.HttpClient(host=chroma_host, port=chroma_port)

        # 임베딩 모델 로드 (더 강력한 모델 사용)
        logger.info("Loading embedding model...")
        # 다국어 성능이 더 좋은 모델로 업그레이드
        self.embedding_model = SentenceTransformer('intfloat/multilingual-e5-large')
        logger.info("Embedding model loaded successfully")

        # MinerU 파서 및 청커 초기화
        logger.info("Initializing MinerU document parser...")

        # 환경변수로 모델 사용 여부 제어 (기본값: True - 실제 모델 사용)
        use_mineru_model = os.getenv("USE_MINERU_MODEL", "true").lower() == "true"

        if use_mineru_model:
            logger.info("Loading MinerU2.5 model for advanced document parsing...")
            self.parser = MinerUDocumentParser(
                use_mock=False,  # 실제 모델 사용
                device="cpu"
            )
        else:
            logger.info("Using heuristic-based document parsing (mock mode)...")
            self.parser = MinerUDocumentParser(use_mock=True)

        self.chunker = LayoutAwareChunker(max_chunk_size=800, overlap=100)
        logger.info("Document parser initialized")

        # 컬렉션 생성/조회
        self.collection = self.chroma_client.get_or_create_collection(
            name="pms_documents_v2",
            metadata={"description": "PMS 프로젝트 산출물 및 문서 (MinerU 기반 구조 파싱)"}
        )

        logger.info(f"Collection initialized: {self.collection.name}")

    def reset_collection(self):
        """컬렉션 초기화 (모든 데이터 삭제)"""
        try:
            logger.warning("Resetting collection - deleting all data...")
            # 기존 컬렉션 삭제
            try:
                self.chroma_client.delete_collection(name="pms_documents_v2")
                logger.info("Deleted existing collection")
            except Exception as e:
                logger.info(f"No existing collection to delete: {e}")

            # 새 컬렉션 생성
            self.collection = self.chroma_client.get_or_create_collection(
                name="pms_documents_v2",
                metadata={"description": "PMS 프로젝트 산출물 및 문서 (MinerU 기반 구조 파싱)"}
            )
            logger.info("Created new collection")
            return True
        except Exception as e:
            logger.error(f"Failed to reset collection: {e}")
            return False

    def extract_text_from_file(self, file_path: str, file_type: str) -> str:
        """파일에서 텍스트 추출 (기존과 동일)"""
        try:
            if file_type.lower() == 'pdf':
                return self._extract_from_pdf(file_path)
            elif file_type.lower() in ['doc', 'docx']:
                return self._extract_from_docx(file_path)
            elif file_type.lower() in ['xls', 'xlsx']:
                return self._extract_from_excel(file_path)
            elif file_type.lower() in ['txt', 'md']:
                with open(file_path, 'r', encoding='utf-8') as f:
                    return f.read()
            else:
                logger.warning(f"Unsupported file type: {file_type}")
                return ""
        except Exception as e:
            logger.error(f"Failed to extract text from {file_path}: {e}")
            return ""

    def _extract_from_pdf(self, file_path: str) -> str:
        """PDF에서 텍스트 추출"""
        reader = PdfReader(file_path)
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return "\n\n".join(text_parts)

    def _extract_from_docx(self, file_path: str) -> str:
        """Word 문서에서 텍스트 추출"""
        doc = Document(file_path)
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        return "\n\n".join(text_parts)

    def _extract_from_excel(self, file_path: str) -> str:
        """Excel에서 텍스트 추출"""
        workbook = openpyxl.load_workbook(file_path, read_only=True)
        text_parts = []

        for sheet in workbook.worksheets:
            text_parts.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
                if row_text.strip():
                    text_parts.append(row_text)

        return "\n".join(text_parts)

    def add_document(self, document: Dict[str, str]) -> bool:
        """
        단일 문서를 벡터 DB에 추가 (MinerU 기반 구조 파싱)

        개선사항:
        1. 문서를 구조적으로 파싱 (제목, 문단, 표 등 인식)
        2. 의미 단위로 청킹 (제목 + 관련 문단)
        3. 구조 정보를 메타데이터에 포함
        """
        try:
            doc_id = document.get('id')
            content = document.get('content', '')
            metadata = document.get('metadata', {})

            if not doc_id or not content:
                logger.error("Document must have 'id' and 'content'")
                return False

            # MinerU 기반 구조 파싱 및 청킹
            logger.info(f"Parsing document {doc_id} with MinerU...")
            blocks = self.parser.parse_document(content, metadata)
            chunks = self.chunker.chunk_blocks(blocks)

            logger.info(f"Document {doc_id} parsed into {len(blocks)} blocks and {len(chunks)} chunks")

            # 각 청크에 대해 임베딩 생성 및 저장
            for i, chunk_data in enumerate(chunks):
                chunk_id = f"{doc_id}_chunk_{i}"
                chunk_content = chunk_data['content']
                chunk_metadata = chunk_data['metadata']

                # E5 모델은 쿼리 프리픽스 사용
                # 저장 시에는 "passage:" 프리픽스
                embedding_text = f"passage: {chunk_content}"
                embedding = self.embedding_model.encode([embedding_text])[0].tolist()

                # 메타데이터 병합
                full_metadata = {
                    **metadata,
                    **chunk_metadata,
                    'chunk_index': i,
                    'total_chunks': len(chunks),
                    'parent_doc_id': doc_id
                }

                self.collection.add(
                    documents=[chunk_content],
                    embeddings=[embedding],
                    metadatas=[full_metadata],
                    ids=[chunk_id]
                )

            logger.info(f"✅ Added document {doc_id} with {len(chunks)} layout-aware chunks")
            return True

        except Exception as e:
            logger.error(f"Failed to add document: {e}", exc_info=True)
            return False

    def add_documents(self, documents: List[Dict[str, str]]) -> int:
        """여러 문서를 벡터 DB에 추가"""
        success_count = 0
        for doc in documents:
            if self.add_document(doc):
                success_count += 1
        return success_count

    def search(self, query: str, top_k: int = 3, filter_metadata: Optional[Dict] = None) -> List[Dict]:
        """
        질의에 대해 관련 문서 검색 (개선된 검색 로직)

        개선사항:
        1. E5 모델의 "query:" 프리픽스 사용
        2. 구조 정보를 활용한 리랭킹
        3. 더 나은 중복 제거
        """
        try:
            # E5 모델은 검색 시 "query:" 프리픽스 사용
            query_with_prefix = f"query: {query}"
            query_embedding = self.embedding_model.encode([query_with_prefix])[0].tolist()

            # 검색 파라미터 (더 많이 검색 후 리랭킹)
            search_params = {
                'query_embeddings': [query_embedding],
                'n_results': min(top_k * 3, 20)  # 중복 제거 및 리랭킹을 위해 더 많이 검색
            }

            if filter_metadata:
                search_params['where'] = filter_metadata

            results = self.collection.query(**search_params)

            # 결과 포맷팅 및 중복 제거
            documents = []
            seen_parent_ids = set()

            if results['documents'] and len(results['documents']) > 0:
                for i, doc in enumerate(results['documents'][0]):
                    metadata = results['metadatas'][0][i] if results['metadatas'] else {}
                    parent_id = metadata.get('parent_doc_id', f'doc_{i}')
                    distance = results['distances'][0][i] if results.get('distances') else 1.0

                    # 같은 문서의 다른 청크는 건너뛰기
                    # 단, 표나 구조화된 데이터는 별도로 포함
                    if parent_id in seen_parent_ids:
                        # 구조화된 데이터는 중복 허용
                        if not metadata.get('is_structured', False):
                            continue

                    seen_parent_ids.add(parent_id)

                    # 구조 정보 기반 점수 보정
                    relevance_score = self._calculate_relevance_score(metadata, distance)

                    documents.append({
                        'content': doc,
                        'metadata': metadata,
                        'distance': distance,
                        'relevance_score': relevance_score
                    })

                # relevance_score 기준 정렬
                documents.sort(key=lambda x: x['relevance_score'], reverse=True)

                # top_k 개수만큼만 반환
                documents = documents[:top_k]

            logger.info(f"Found {len(documents)} relevant documents for query: {query[:50]}...")
            return documents

        except Exception as e:
            logger.error(f"Search failed: {e}", exc_info=True)
            return []

    def _calculate_relevance_score(self, metadata: Dict, distance: float) -> float:
        """
        구조 정보를 활용한 관련성 점수 계산

        낮은 distance = 높은 유사도
        구조화된 데이터(표, 리스트) = 가중치 부여
        제목 포함 = 가중치 부여
        """
        # 거리를 유사도로 변환 (0~1)
        similarity = max(0, 1 - distance)

        # 구조 가중치
        structure_weight = 1.0
        if metadata.get('has_table', False):
            structure_weight = 1.2  # 표 포함 시 20% 가중
        elif metadata.get('has_list', False):
            structure_weight = 1.1  # 리스트 포함 시 10% 가중

        # 제목 가중치
        title_weight = 1.1 if metadata.get('title') else 1.0

        # 최종 점수
        relevance_score = similarity * structure_weight * title_weight

        return relevance_score

    def delete_document(self, doc_id: str) -> bool:
        """문서 삭제 (모든 청크 포함)"""
        try:
            # 해당 문서의 모든 청크 찾기
            results = self.collection.get(
                where={"parent_doc_id": doc_id}
            )

            if results['ids']:
                self.collection.delete(ids=results['ids'])
                logger.info(f"Deleted document {doc_id} with {len(results['ids'])} chunks")
                return True
            else:
                logger.warning(f"No chunks found for document {doc_id}")
                return False

        except Exception as e:
            logger.error(f"Failed to delete document {doc_id}: {e}")
            return False

    def get_collection_stats(self) -> Dict:
        """컬렉션 통계 조회"""
        try:
            count = self.collection.count()

            # 구조 타입별 통계
            all_items = self.collection.get()
            structured_count = 0
            table_count = 0
            list_count = 0

            if all_items['metadatas']:
                for metadata in all_items['metadatas']:
                    if metadata.get('is_structured', False):
                        structured_count += 1
                    if metadata.get('has_table', False):
                        table_count += 1
                    if metadata.get('has_list', False):
                        list_count += 1

            return {
                'total_chunks': count,
                'collection_name': self.collection.name,
                'structured_chunks': structured_count,
                'chunks_with_tables': table_count,
                'chunks_with_lists': list_count,
                'parser': 'MinerU2.5-based'
            }
        except Exception as e:
            logger.error(f"Failed to get collection stats: {e}")
            return {'error': str(e)}
