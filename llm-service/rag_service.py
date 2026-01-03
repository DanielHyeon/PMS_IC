"""
RAG (Retrieval-Augmented Generation) 서비스
ChromaDB를 사용한 문서 임베딩 및 검색
"""

import chromadb
from sentence_transformers import SentenceTransformer
from typing import List, Dict, Optional
import logging
import os
from pypdf import PdfReader
from docx import Document
import openpyxl

logger = logging.getLogger(__name__)

class RAGService:
    def __init__(self):
        # ChromaDB 클라이언트 초기화
        chroma_host = os.getenv("CHROMA_HOST", "localhost")
        chroma_port = int(os.getenv("CHROMA_PORT", "8000"))

        logger.info(f"Connecting to ChromaDB at {chroma_host}:{chroma_port}")
        self.chroma_client = chromadb.HttpClient(host=chroma_host, port=chroma_port)

        # 임베딩 모델 로드 (다국어 지원)
        logger.info("Loading embedding model...")
        self.embedding_model = SentenceTransformer('paraphrase-multilingual-MiniLM-L12-v2')
        logger.info("Embedding model loaded successfully")

        # 컬렉션 생성/조회
        self.collection = self.chroma_client.get_or_create_collection(
            name="pms_documents",
            metadata={"description": "PMS 프로젝트 산출물 및 문서"}
        )

        # 초기 데이터 시딩 (컬렉션이 비어 있을 때만)
        try:
            if self.collection.count() == 0:
                from init_rag_mock_data import MOCK_DOCUMENTS
                logger.info("Seeding mock documents into ChromaDB collection")
                self.add_documents(MOCK_DOCUMENTS)
        except Exception as e:
            logger.error(f"Failed to seed mock documents: {e}", exc_info=True)

    def extract_text_from_file(self, file_path: str, file_type: str) -> str:
        """파일에서 텍스트 추출"""
        try:
            if file_type.lower() == 'pdf':
                return self._extract_from_pdf(file_path)
            elif file_type.lower() in ['doc', 'docx']:
                return self._extract_from_docx(file_path)
            elif file_type.lower() in ['xls', 'xlsx']:
                return self._extract_from_excel(file_path)
            elif file_type.lower() in ['txt', 'md']:
                with open(file_path, 'r', encoding='utf-8') as f:
                    return f.read()
            else:
                logger.warning(f"Unsupported file type: {file_type}")
                return ""
        except Exception as e:
            logger.error(f"Failed to extract text from {file_path}: {e}")
            return ""

    def _extract_from_pdf(self, file_path: str) -> str:
        """PDF에서 텍스트 추출"""
        reader = PdfReader(file_path)
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return "\n\n".join(text_parts)

    def _extract_from_docx(self, file_path: str) -> str:
        """Word 문서에서 텍스트 추출"""
        doc = Document(file_path)
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        return "\n\n".join(text_parts)

    def _extract_from_excel(self, file_path: str) -> str:
        """Excel에서 텍스트 추출"""
        workbook = openpyxl.load_workbook(file_path, read_only=True)
        text_parts = []

        for sheet in workbook.worksheets:
            text_parts.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
                if row_text.strip():
                    text_parts.append(row_text)

        return "\n".join(text_parts)

    def chunk_text(self, text: str, chunk_size: int = 512, overlap: int = 50) -> List[str]:
        """텍스트를 청크로 분할"""
        words = text.split()
        chunks = []

        for i in range(0, len(words), chunk_size - overlap):
            chunk = " ".join(words[i:i + chunk_size])
            if chunk.strip():
                chunks.append(chunk)

        return chunks

    def add_document(self, document: Dict[str, str]) -> bool:
        """단일 문서를 벡터 DB에 추가"""
        try:
            doc_id = document.get('id')
            content = document.get('content', '')
            metadata = document.get('metadata', {})

            if not doc_id or not content:
                logger.error("Document must have 'id' and 'content'")
                return False

            # 텍스트를 청크로 분할
            chunks = self.chunk_text(content)

            # 각 청크에 대해 임베딩 생성 및 저장
            for i, chunk in enumerate(chunks):
                chunk_id = f"{doc_id}_chunk_{i}"
                embedding = self.embedding_model.encode([chunk])[0].tolist()

                chunk_metadata = {
                    **metadata,
                    'chunk_index': i,
                    'total_chunks': len(chunks),
                    'parent_doc_id': doc_id
                }

                self.collection.add(
                    documents=[chunk],
                    embeddings=[embedding],
                    metadatas=[chunk_metadata],
                    ids=[chunk_id]
                )

            logger.info(f"Added document {doc_id} with {len(chunks)} chunks")
            return True

        except Exception as e:
            logger.error(f"Failed to add document: {e}", exc_info=True)
            return False

    def add_documents(self, documents: List[Dict[str, str]]) -> int:
        """여러 문서를 벡터 DB에 추가"""
        success_count = 0
        for doc in documents:
            if self.add_document(doc):
                success_count += 1
        return success_count

    def search(self, query: str, top_k: int = 3, filter_metadata: Optional[Dict] = None) -> List[Dict]:
        """질의에 대해 관련 문서 검색"""
        try:
            query_embedding = self.embedding_model.encode([query])[0].tolist()

            # 검색 파라미터
            search_params = {
                'query_embeddings': [query_embedding],
                'n_results': top_k * 2  # 중복 제거를 위해 더 많이 검색
            }

            if filter_metadata:
                search_params['where'] = filter_metadata

            results = self.collection.query(**search_params)

            # 결과 포맷팅 및 중복 제거
            documents = []
            seen_parent_ids = set()

            if results['documents'] and len(results['documents']) > 0:
                for i, doc in enumerate(results['documents'][0]):
                    metadata = results['metadatas'][0][i] if results['metadatas'] else {}
                    parent_id = metadata.get('parent_doc_id', f'doc_{i}')

                    # 같은 문서의 다른 청크는 건너뛰기
                    if parent_id in seen_parent_ids:
                        continue

                    seen_parent_ids.add(parent_id)

                    documents.append({
                        'content': doc,
                        'metadata': metadata,
                        'distance': results['distances'][0][i] if results.get('distances') else None
                    })

                    # top_k 개수만큼만 반환
                    if len(documents) >= top_k:
                        break

            logger.info(f"Found {len(documents)} relevant documents for query: {query[:50]}...")
            return documents

        except Exception as e:
            logger.error(f"Search failed: {e}", exc_info=True)
            return []

    def delete_document(self, doc_id: str) -> bool:
        """문서 삭제 (모든 청크 포함)"""
        try:
            # 해당 문서의 모든 청크 찾기
            results = self.collection.get(
                where={"parent_doc_id": doc_id}
            )

            if results['ids']:
                self.collection.delete(ids=results['ids'])
                logger.info(f"Deleted document {doc_id} with {len(results['ids'])} chunks")
                return True
            else:
                logger.warning(f"No chunks found for document {doc_id}")
                return False

        except Exception as e:
            logger.error(f"Failed to delete document {doc_id}: {e}")
            return False

    def get_collection_stats(self) -> Dict:
        """컬렉션 통계 조회"""
        try:
            count = self.collection.count()
            return {
                'total_chunks': count,
                'collection_name': self.collection.name
            }
        except Exception as e:
            logger.error(f"Failed to get collection stats: {e}")
            return {'error': str(e)}
